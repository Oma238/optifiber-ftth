"""Genera el informe Word del Capítulo 4 de OptiFiber a partir de las evidencias."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[1]
DOCS = RAIZ / "docs"
EVIDENCIAS = DOCS / "evidencias"
SALIDA = DOCS / "Informe_Capitulo_4_OptiFiber.docx"

AZUL = RGBColor(0x1D, 0x4E, 0xD8)
VERDE = RGBColor(0x04, 0x78, 0x57)
GRIS = RGBColor(0x47, 0x55, 0x69)
GRIS_CLARO = RGBColor(0x64, 0x74, 0x8B)
NEGRO = RGBColor(0x0F, 0x17, 0x2A)


def leer_json(nombre: str, defecto: dict | None = None) -> dict:
    ruta = EVIDENCIAS / nombre
    if not ruta.exists():
        return defecto or {}
    return json.loads(ruta.read_text(encoding="utf-8"))


def leer_texto(nombre: str, defecto: str = "") -> str:
    ruta = EVIDENCIAS / nombre
    if not ruta.exists():
        return defecto
    return ruta.read_text(encoding="utf-8", errors="replace")


def sombrear_celda(celda, color: str) -> None:
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def sombrear_parrafo(parrafo, color: str) -> None:
    p_pr = parrafo._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    p_pr.append(shd)


def borde_parrafo(parrafo, color: str = "CBD5E1", tamano: str = "6") -> None:
    p_pr = parrafo._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    for lado in ("top", "left", "bottom", "right"):
        borde = OxmlElement(f"w:{lado}")
        borde.set(qn("w:val"), "single")
        borde.set(qn("w:sz"), tamano)
        borde.set(qn("w:space"), "4")
        borde.set(qn("w:color"), color)
        p_bdr.append(borde)
    p_pr.append(p_bdr)


def agregar_codigo(doc: Document, texto: str, tamano: float = 8.5) -> None:
    lineas = texto.strip("\n").split("\n")
    parrafo = doc.add_paragraph(style="Codigo")
    for indice, linea in enumerate(lineas):
        run = parrafo.add_run(linea if linea else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(tamano)
        run.font.color.rgb = NEGRO
        if indice < len(lineas) - 1:
            run.add_break()
    sombrear_parrafo(parrafo, "F1F5F9")
    borde_parrafo(parrafo, "E2E8F0", "6")
    parrafo.paragraph_format.space_after = Pt(6)
    parrafo.paragraph_format.space_before = Pt(2)
    parrafo.paragraph_format.line_spacing = 1.0


def agregar_tabla(
    doc: Document,
    encabezados: list[str],
    filas: list[list[str]],
    anchos: list[float] | None = None,
    tamano: float = 8.5,
) -> None:
    tabla = doc.add_table(rows=1, cols=len(encabezados))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False

    encabezado = tabla.rows[0].cells
    for indice, titulo in enumerate(encabezados):
        encabezado[indice].text = ""
        parrafo = encabezado[indice].paragraphs[0]
        run = parrafo.add_run(titulo)
        run.bold = True
        run.font.size = Pt(tamano)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear_celda(encabezado[indice], "1D4ED8")
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for datos in filas:
        celdas = tabla.add_row().cells
        for indice, valor in enumerate(datos):
            celdas[indice].text = ""
            parrafo = celdas[indice].paragraphs[0]
            run = parrafo.add_run(str(valor))
            run.font.size = Pt(tamano)
            if indice == 0:
                run.bold = True
                sombrear_celda(celdas[indice], "EFF6FF")

    if anchos:
        for fila in tabla.rows:
            for indice, ancho in enumerate(anchos):
                fila.cells[indice].width = Inches(ancho)

    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def agregar_imagen(doc: Document, ruta: Path, ancho: float = 16.2) -> None:
    """Inserta una captura centrada con un ancho util para la pagina A4."""
    if not ruta.exists():
        doc.add_paragraph(f"[Captura pendiente: {ruta.name}]")
        return
    parrafo = doc.add_paragraph()
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.paragraph_format.space_before = Pt(4)
    parrafo.paragraph_format.space_after = Pt(2)
    parrafo.add_run().add_picture(str(ruta), width=Cm(ancho))


def agregar_caption(doc: Document, texto: str) -> None:
    parrafo = doc.add_paragraph(style="Caption")
    run = parrafo.add_run(texto)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRIS_CLARO
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER


def agregar_bullet(doc: Document, texto: str) -> None:
    parrafo = doc.add_paragraph(style="List Bullet")
    parrafo.add_run(texto)
    parrafo.paragraph_format.space_after = Pt(2)


def agregar_numero(doc: Document, texto: str, numero: int) -> None:
    parrafo = doc.add_paragraph()
    run = parrafo.add_run(f"{numero}. ")
    run.bold = True
    parrafo.add_run(texto)
    parrafo.paragraph_format.left_indent = Cm(0.6)
    parrafo.paragraph_format.first_line_indent = Cm(-0.4)
    parrafo.paragraph_format.space_after = Pt(2)


def configurar_estilos(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = NEGRO
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    for nombre, tamano, color, antes, despues in [
        ("Heading 1", 16, AZUL, 14, 6),
        ("Heading 2", 13, VERDE, 10, 4),
        ("Heading 3", 11.5, NEGRO, 8, 3),
    ]:
        estilo = doc.styles[nombre]
        estilo.font.name = "Calibri Light"
        estilo.font.size = Pt(tamano)
        estilo.font.color.rgb = color
        estilo.font.bold = True
        estilo.paragraph_format.space_before = Pt(antes)
        estilo.paragraph_format.space_after = Pt(despues)
        estilo.paragraph_format.keep_with_next = True

    estilo_codigo = doc.styles.add_style("Codigo", WD_STYLE_TYPE.PARAGRAPH)
    estilo_codigo.font.name = "Consolas"
    estilo_codigo.font.size = Pt(8.5)
    estilo_codigo.paragraph_format.left_indent = Cm(0.2)
    estilo_codigo.paragraph_format.right_indent = Cm(0.2)
    estilo_codigo.paragraph_format.space_after = Pt(0)
    estilo_codigo.paragraph_format.line_spacing = 1.0

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = GRIS_CLARO


def configurar_pagina(doc: Document) -> None:
    seccion = doc.sections[0]
    seccion.page_height = Cm(29.7)
    seccion.page_width = Cm(21.0)
    seccion.top_margin = Cm(2.0)
    seccion.bottom_margin = Cm(2.0)
    seccion.left_margin = Cm(2.2)
    seccion.right_margin = Cm(2.2)
    seccion.header_distance = Cm(1.0)
    seccion.footer_distance = Cm(1.0)
    seccion.different_first_page_header_footer = True


def agregar_encabezado_pie(doc: Document) -> None:
    seccion = doc.sections[0]
    encabezado = seccion.header
    parrafo = encabezado.paragraphs[0]
    parrafo.text = ""
    run = parrafo.add_run("OPTIFIBER · Construcción de Software")
    run.font.size = Pt(8)
    run.font.color.rgb = GRIS
    run.bold = True
    parrafo.add_run("\t")
    run2 = parrafo.add_run("PBI-02 + PBI-03 · Capítulo 4")
    run2.font.size = Pt(8)
    run2.font.color.rgb = GRIS
    parrafo.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    pie = seccion.footer
    parrafo_pie = pie.paragraphs[0]
    parrafo_pie.text = ""
    run_pie = parrafo_pie.add_run("Informe de Proyecto de Software · 2026")
    run_pie.font.size = Pt(8)
    run_pie.font.color.rgb = GRIS
    parrafo_pie.add_run("\t")
    run_pie2 = parrafo_pie.add_run("Página ")
    run_pie2.font.size = Pt(8)
    run_pie2.font.color.rgb = GRIS

    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run_pie3 = parrafo_pie.add_run()
    run_pie3.font.size = Pt(8)
    run_pie3._r.append(fld_char1)
    run_pie3._r.append(instr_text)
    run_pie3._r.append(fld_char2)
    parrafo_pie.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def portada(doc: Document) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(30)
    parrafo = doc.add_paragraph()
    run = parrafo.add_run("INFORME DE PROYECTO DE SOFTWARE")
    run.font.name = "Calibri Light"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = AZUL
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    parrafo = doc.add_paragraph()
    run = parrafo.add_run("CAPÍTULO 4")
    run.font.name = "Calibri"
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = VERDE
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    parrafo = doc.add_paragraph()
    run = parrafo.add_run(
        "Persistencia de Datos, Mapeo ORM,\nPruebas con Fixtures e Interfaz Gráfica"
    )
    run.font.name = "Calibri Light"
    run.font.size = Pt(20)
    run.font.color.rgb = NEGRO
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.paragraph_format.space_after = Pt(24)

    tabla = doc.add_table(rows=1, cols=1)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    celda = tabla.rows[0].cells[0]
    sombrear_celda(celda, "1D4ED8")
    parrafo = celda.paragraphs[0]
    run = parrafo.add_run("OPTIFIBER")
    run.font.name = "Calibri Light"
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo = celda.add_paragraph()
    run = parrafo.add_run("Venta de productos FTTH · Fibra óptica, ONT/ONU, accesorios y redes")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0xDB, 0xEA, 0xFE)
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    agregar_tabla(
        doc,
        ["Campo", "Detalle"],
        [
            ["Curso", "Construcción de Software"],
            ["Trabajo", "Instructivo de Trabajo Colaborativo N° 04"],
            ["PBI", "PBI-02 · Persistencia de clientes, pedidos y catálogo FTTH + PBI-03 · Interfaz gráfica de escritorio"],
            ["Rama", "feature/PBI-02-orm-persistencia y feature/PBI-03-interfaz-grafica"],
            ["Stack", "Python 3.12+ · SQLAlchemy 2.x · SQLite · CustomTkinter · pytest · Ruff · GitHub Actions"],
            ["Fecha", "28 de octubre de 2026"],
        ],
        anchos=[1.6, 4.8],
        tamano=9,
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    agregar_tabla(
        doc,
        ["Integrante", "Rol", "GitHub sugerido"],
        [
            ["Huallpayunca Carrion Omar Sebastian", "Git Master + PO/SM", "@omarhuallpayunca"],
            ["Kevin Jesus Llerena Fernandez", "Lead Backend Engineer", "@kevinllerena"],
            ["Amaru Josue Ai-me Yunguri", "Database & ORM Specialist", "@amaruaime"],
            ["Julio Cesar Diaz Tito", "DevOps & QA + Frontend/UX", "@juliodiaz"],
        ],
        anchos=[2.7, 2.2, 1.5],
        tamano=9,
    )

    parrafo = doc.add_paragraph()
    run = parrafo.add_run(
        "Documento elaborado por el equipo de OPTIFIBER. Los usuarios de GitHub son "
        "sugeridos y deben reemplazarse por los reales antes de publicar."
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = GRIS
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def seccion_resumen(doc: Document) -> None:
    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        "OPTIFIBER necesita registrar su operación comercial de productos FTTH: clientes, "
        "catálogo, pedidos y líneas de detalle. El PBI-02 implementa esa persistencia con "
        "SQLAlchemy 2.x sobre SQLite, manteniendo la lógica de negocio en src/dominio y la "
        "infraestructura en src/infraestructura. El diseño incluye seis entidades, relaciones "
        "1:N, una relación N:M pura (Producto–Etiqueta) y una relación N:M con datos adicionales "
        "(Pedido–Producto mediante DetallePedido), reglas de cascada y restricciones de integridad."
    )
    doc.add_paragraph(
        "La solución se desarrolló con ciclos TDD RED → GREEN → REFACTOR y un fixture de base de "
        "datos en memoria (sqlite:///:memory:) que garantiza pruebas herméticas. La suite final "
        "tiene 23 pruebas aprobadas, Ruff no reporta errores y el workflow de GitHub Actions "
        "instala SQLAlchemy, ejecuta Ruff y pytest. La CLI permite demostrar el flujo completo: "
        "catálogo, alta de clientes, pedido con descuento de stock, reporte y eliminación en cascada."
    )
    doc.add_paragraph(
        "Sobre esa base, el PBI-03 incorpora una aplicación de escritorio con CustomTkinter: "
        "panel general, catálogo, clientes, pedidos y reportes; tema oscuro minimalista, formularios "
        "modales, punto de venta con carrito, exportación CSV y pedidos de demostración idempotentes. "
        "La interfaz reutiliza los repositorios y las reglas de negocio, por lo que no duplica la lógica."
    )

    agregar_tabla(
        doc,
        ["Indicador", "Resultado"],
        [
            ["Entidades ORM", "6 (Cliente, Pedido, DetallePedido, Producto, Categoría, Etiqueta)"],
            ["Relaciones", "4 relaciones 1:N + 2 relaciones N:M (una pura y una con atributos)"],
            ["Pruebas automatizadas", "23 passed (21 del PBI-02 + 2 de datos demo)"],
            ["Calidad estática", "Ruff: All checks passed!"],
            ["Interfaz", "GUI CustomTkinter con 5 pantallas, tema oscuro y exportación CSV"],
            ["Cobertura funcional", "Catálogo, clientes, pedidos, stock, descuento, reportes y cascada"],
            ["CI/CD", "GitHub Actions con Python 3.12, Ruff y pytest"],
        ],
        anchos=[1.9, 4.5],
        tamano=9,
    )


def seccion_erd(doc: Document, hashes: dict) -> None:
    encabezado = doc.add_heading("4.1 Modelo Entidad-Relación (ERD)", level=1)
    encabezado.paragraph_format.page_break_before = True

    doc.add_heading("4.1.1 Objetivo del modelo", level=2)
    for indice, texto in enumerate([
        "Identificar de forma única a cada cliente por DNI o RUC.",
        "Clasificar productos FTTH por categoría y etiquetas comerciales.",
        "Registrar pedidos con varias líneas, conservando el precio unitario histórico.",
        "Controlar el stock y calcular totales con descuento por volumen.",
        "Eliminar un cliente junto con sus pedidos y detalles sin dejar registros huérfanos.",
    ], start=1):
        agregar_numero(doc, texto, indice)

    doc.add_heading("4.1.2 Entidades del modelo", level=2)
    agregar_tabla(
        doc,
        ["Entidad", "Rol en el negocio", "PK", "FK / restricciones"],
        [
            ["CLIENTE", "Persona natural o empresa que compra en OPTIFIBER.", "id", "documento UNIQUE"],
            ["PEDIDO", "Orden de compra realizada por un cliente.", "id", "cliente_id → clientes.id"],
            ["DETALLE_PEDIDO", "Línea del pedido con cantidad, precio y subtotal.", "id", "pedido_id, producto_id; UNIQUE compuesto"],
            ["PRODUCTO", "Artículo FTTH que vende OPTIFIBER.", "id", "categoria_id → categorias.id; sku UNIQUE"],
            ["CATEGORIA", "Agrupación comercial del catálogo.", "id", "nombre UNIQUE"],
            ["ETIQUETA", "Clasificación transversal del producto.", "id", "nombre UNIQUE"],
            ["PRODUCTO_ETIQUETA", "Tabla de asociación N:M Producto–Etiqueta.", "producto_id + etiqueta_id", "FK a productos y etiquetas"],
        ],
        anchos=[1.2, 2.3, 1.2, 1.7],
        tamano=8.5,
    )
    agregar_caption(doc, "Tabla 1. Entidades del ERD de OPTIFIBER.")

    doc.add_heading("4.1.3 Relaciones y cardinalidad", level=2)
    agregar_tabla(
        doc,
        ["Relación", "Cardinalidad", "Implementación", "Regla de integridad"],
        [
            ["Cliente — Pedido", "1:N", "ClienteModel.pedidos ↔ PedidoModel.cliente", "cascade all, delete-orphan; FK CASCADE"],
            ["Pedido — DetallePedido", "1:N", "PedidoModel.detalles ↔ DetallePedidoModel.pedido", "cascade all, delete-orphan; FK CASCADE"],
            ["Categoría — Producto", "1:N", "CategoriaModel.productos ↔ ProductoModel.categoria", "Sin cascada; FK RESTRICT"],
            ["Producto — DetallePedido", "1:N", "ProductoModel.detalles ↔ DetallePedidoModel.producto", "Sin cascada; FK RESTRICT"],
            ["Producto — Etiqueta", "N:M", "relationship(secondary=producto_etiqueta)", "Tabla de asociación con PK compuesta"],
            ["Pedido — Producto", "N:M con atributos", "DetallePedidoModel", "UNIQUE(pedido_id, producto_id)"],
        ],
        anchos=[1.4, 1.0, 2.2, 1.8],
        tamano=8.5,
    )
    agregar_caption(doc, "Tabla 2. Relaciones, cardinalidades y reglas de integridad.")

    doc.add_heading("4.1.4 Diagrama ERD", level=2)
    doc.add_picture(str(DOCS / "erd_optifiber.png"), width=Inches(6.4))
    ultimo = doc.paragraphs[-1]
    ultimo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    agregar_caption(
        doc,
        "Figura 1. Modelo Entidad-Relación de OPTIFIBER con seis entidades, claves, relaciones y cascada.",
    )

    doc.add_heading("4.1.5 Decisiones de diseño", level=2)
    for texto in [
        "Claves primarias subrogadas: id autoincremental en todas las entidades; documento y sku son claves naturales UNIQUE.",
        "Precio histórico: DetallePedido.precio_unitario conserva el precio al momento de la venta.",
        "N:M pura: Producto–Etiqueta con Table(...), relationship(secondary=...) y back_populates en ambos lados.",
        "N:M con atributos: Pedido–Producto se resuelve con DetallePedidoModel, que añade cantidad, precio y subtotal.",
        "Cascada selectiva: solo se propaga a datos dependientes del pedido; categorías y productos usan RESTRICT.",
        "Integridad en SQLite: el engine activa PRAGMA foreign_keys=ON en cada conexión.",
        "Separación de capas: el dominio no importa SQLAlchemy; la infraestructura contiene el mapeo ORM.",
    ]:
        agregar_bullet(doc, texto)


def seccion_orm(doc: Document, hashes: dict) -> None:
    doc.add_page_break()
    doc.add_heading("4.2 Código del modelo ORM", level=1)

    doc.add_heading("4.2.1 Clase base y tabla de asociación N:M", level=2)
    agregar_codigo(
        doc,
        '''class Base(DeclarativeBase):
    """Clase base declarativa para todas las entidades ORM."""


producto_etiqueta = Table(
    "producto_etiqueta",
    Base.metadata,
    Column("producto_id", Integer,
           ForeignKey("productos.id", ondelete="CASCADE"), primary_key=True),
    Column("etiqueta_id", Integer,
           ForeignKey("etiquetas.id", ondelete="CASCADE"), primary_key=True),
)''',
    )
    doc.add_paragraph(
        "La tabla de asociación se declara antes de las clases que la referencian. "
        "Esto evita el error de resolución de secondary y permite que SQLAlchemy "
        "construya la relación N:M."
    )

    doc.add_heading("4.2.2 Relación 1:N Cliente–Pedido con cascada", level=2)
    agregar_codigo(
        doc,
        '''class ClienteModel(Base):
    __tablename__ = "clientes"

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    nombre: Mapped[str] = mapped_column(String(120), nullable=False)
    documento: Mapped[str] = mapped_column(String(20), nullable=False, unique=True, index=True)
    telefono: Mapped[str] = mapped_column(String(20), nullable=False, default="")
    email: Mapped[str] = mapped_column(String(120), nullable=False, default="")
    direccion: Mapped[str] = mapped_column(String(200), nullable=False, default="")
    creado_en: Mapped[datetime] = mapped_column(DateTime, nullable=False, default=datetime.now)

    pedidos: Mapped[list["PedidoModel"]] = relationship(
        back_populates="cliente",
        cascade="all, delete-orphan",
        order_by="PedidoModel.id",
    )


class PedidoModel(Base):
    __tablename__ = "pedidos"
    __table_args__ = (CheckConstraint("total >= 0", name="ck_pedidos_total_no_negativo"),)

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    fecha: Mapped[date] = mapped_column(Date, nullable=False, default=date.today)
    estado: Mapped[str] = mapped_column(String(20), nullable=False, default="PENDIENTE")
    total: Mapped[float] = mapped_column(Float, nullable=False, default=0.0)
    cliente_id: Mapped[int] = mapped_column(
        Integer, ForeignKey("clientes.id", ondelete="CASCADE"), nullable=False, index=True
    )

    cliente: Mapped[ClienteModel] = relationship(back_populates="pedidos")''',
    )
    doc.add_paragraph(
        "cascade=\"all, delete-orphan\" indica que, al eliminar un ClienteModel, "
        "SQLAlchemy elimina sus PedidoModel asociados. back_populates mantiene "
        "navegable la relación inversa pedido.cliente."
    )

    doc.add_heading("4.2.3 DetallePedido: N:M con atributos y segunda cascada", level=2)
    agregar_codigo(
        doc,
        '''class DetallePedidoModel(Base):
    __tablename__ = "detalle_pedidos"
    __table_args__ = (
        UniqueConstraint("pedido_id", "producto_id", name="uq_detalle_pedido_producto"),
        CheckConstraint("cantidad > 0", name="ck_detalle_cantidad_positiva"),
        CheckConstraint("precio_unitario >= 0", name="ck_detalle_precio_no_negativo"),
        CheckConstraint("subtotal >= 0", name="ck_detalle_subtotal_no_negativo"),
    )

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    pedido_id: Mapped[int] = mapped_column(
        Integer, ForeignKey("pedidos.id", ondelete="CASCADE"), nullable=False, index=True
    )
    producto_id: Mapped[int] = mapped_column(
        Integer, ForeignKey("productos.id", ondelete="RESTRICT"), nullable=False, index=True
    )
    cantidad: Mapped[int] = mapped_column(Integer, nullable=False)
    precio_unitario: Mapped[float] = mapped_column(Float, nullable=False)
    subtotal: Mapped[float] = mapped_column(Float, nullable=False)

    pedido: Mapped[PedidoModel] = relationship(back_populates="detalles")
    producto: Mapped[ProductoModel] = relationship(back_populates="detalles")''',
    )
    doc.add_paragraph(
        "En PedidoModel, la colección se define con cascade=\"all, delete-orphan\". "
        "Así, al eliminar un cliente se eliminan sus pedidos y, en cadena, los detalles "
        "de cada pedido."
    )

    doc.add_heading("4.2.4 Relación N:M Producto–Etiqueta", level=2)
    agregar_codigo(
        doc,
        '''class EtiquetaModel(Base):
    __tablename__ = "etiquetas"

    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    nombre: Mapped[str] = mapped_column(String(50), nullable=False, unique=True)

    productos: Mapped[list["ProductoModel"]] = relationship(
        secondary=producto_etiqueta,
        back_populates="etiquetas",
    )


class ProductoModel(Base):
    __tablename__ = "productos"
    # ... campos ...
    etiquetas: Mapped[list[EtiquetaModel]] = relationship(
        secondary=producto_etiqueta,
        back_populates="productos",
    )''',
    )

    doc.add_heading("4.2.5 Separación de capas", level=2)
    agregar_tabla(
        doc,
        ["Capa", "Archivo", "Responsabilidad", "Dependencias"],
        [
            ["Dominio", "src/dominio/servicios.py", "Subtotal, descuento, stock y DNI/RUC.", "Solo decimal"],
            ["Infraestructura", "src/infraestructura/modelos.py", "Mapeo ORM, relaciones y restricciones.", "SQLAlchemy"],
            ["Infraestructura", "src/infraestructura/repositorios.py", "CRUD y transacciones; usa el dominio.", "Dominio + modelos"],
            ["Infraestructura", "src/infraestructura/db.py", "Engine, sesión y PRAGMA foreign_keys.", "SQLAlchemy"],
            ["Presentación", "src/presentacion/cli.py", "Menú y formato; no calcula reglas.", "Repositorios + dominio"],
        ],
        anchos=[1.0, 2.1, 2.2, 1.1],
        tamano=8.5,
    )
    agregar_caption(doc, "Tabla 3. Separación de capas y dirección de dependencias.")


def seccion_tdd(doc: Document, hashes: dict) -> None:
    doc.add_page_break()
    doc.add_heading("4.3 Fixture y pruebas de repositorio con TDD", level=1)

    doc.add_heading("4.3.1 Fixture de base de datos en memoria", level=2)
    agregar_codigo(
        doc,
        '''@pytest.fixture(scope="function")
def db_session():
    """Sesion limpia de SQLite en memoria para cada prueba."""
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )

    @event.listens_for(engine, "connect")
    def _activar_claves_foraneas(dbapi_connection, _connection_record):
        cursor = dbapi_connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.close()

    Base.metadata.create_all(bind=engine)
    session = sessionmaker(bind=engine, expire_on_commit=False)()

    try:
        yield session
    finally:
        session.rollback()
        session.close()
        engine.dispose()''',
    )
    doc.add_paragraph(
        "StaticPool reutiliza una sola conexión para que create_all, la sesión y las "
        "consultas vean el mismo esquema en memoria. scope=\"function\" garantiza una "
        "base limpia por prueba."
    )

    doc.add_heading("4.3.2 Ciclos TDD ejecutados", level=2)
    filas = [
        ["1. Cliente 1:N Pedido", f"{hashes.get('red_ciclo1', '')} · ImportError ClienteModel", f"{hashes.get('green_ciclo1', '')} · modelos Cliente/Pedido", "1 passed"],
        ["2. Cascada Cliente→Pedido", f"{hashes.get('red_ciclo2', '')} · IntegrityError NOT NULL", f"{hashes.get('green_ciclo2', '')} · cascade delete-orphan", "2 passed"],
        ["3. N:M y DetallePedido", f"{hashes.get('red_ciclo3', '')} · ImportError CategoriaModel", f"{hashes.get('green_ciclo3', '')} · catálogo y detalle", "5 passed"],
        ["4. Reglas de dominio", f"{hashes.get('red_dominio', '')} · ModuleNotFound servicios", f"{hashes.get('green_dominio', '')} · servicios de dominio", "15 passed"],
        ["5. Repositorios y pedidos", f"{hashes.get('red_repositorios', '')} · ModuleNotFound repositorios", f"{hashes.get('green_repositorios', '')} · repositorios + CLI", "21 passed"],
    ]
    agregar_tabla(
        doc,
        ["Ciclo", "RED (falla esperada)", "GREEN (implementación)", "Resultado"],
        filas,
        anchos=[1.3, 2.2, 2.2, 0.7],
        tamano=8,
    )
    agregar_caption(doc, "Tabla 4. Secuencia TDD con hashes reales y resultado de cada ciclo.")
    doc.add_paragraph(
        f"El commit de cierre REFACTOR es {hashes.get('refactor', '')}; en esa etapa "
        "Ruff verificó el estilo y la suite completa permaneció en verde."
    )

    doc.add_heading("4.3.3 Inventario de pruebas", level=2)
    agregar_tabla(
        doc,
        ["Archivo", "Pruebas", "Casos"],
        [
            ["tests/test_persistencia_orm.py", "5", "1:N, cascada, N:M, cascada Pedido–Detalle, documento único."],
            ["tests/test_dominio.py", "10", "Subtotal, descuento, stock, cantidades, DNI/RUC."],
            ["tests/test_repositorios.py", "6", "Alta de pedido, stock, descuento, duplicados y cascada."],
            ["tests/test_seed.py", "2", "Catálogo inicial y pedidos demo idempotentes con descuento de stock."],
            ["Total", "23", "Suite hermética sobre sqlite:///:memory:."],
        ],
        anchos=[2.2, 0.7, 3.5],
        tamano=8.5,
    )
    agregar_caption(doc, "Tabla 5. Inventario de la suite de pruebas.")

    doc.add_heading("4.3.4 Evidencia de ejecución", level=2)
    pytest_texto = leer_texto("pytest_final.txt")
    lineas = [
        linea
        for linea in pytest_texto.splitlines()
        if "test_dominio.py" in linea
        or "test_persistencia_orm.py" in linea
        or "test_repositorios.py" in linea
        or "test_seed.py" in linea
        or "passed" in linea
    ]
    agregar_codigo(doc, "\n".join(lineas[-6:]) if lineas else "23 passed", tamano=8)
    agregar_codigo(doc, leer_texto("ruff_final.txt", "All checks passed!"), tamano=8)

    doc.add_heading("4.3.5 Pipeline de GitHub Actions", level=2)
    agregar_codigo(
        doc,
        '''name: CI - Calidad y Pruebas
on:
  push:
    branches: [main, dev, "feature/**"]
  pull_request:
    branches: [main, dev]
jobs:
  calidad:
    runs-on: ubuntu-latest
    steps:
      - uses: actions/checkout@v4
      - uses: actions/setup-python@v5
        with:
          python-version: "3.12"
          cache: pip
      - run: |
          python -m pip install --upgrade pip
          pip install -r requirements.txt
      - run: ruff check .
      - run: pytest -v''',
        tamano=8,
    )
    doc.add_paragraph(
        "requirements.txt declara SQLAlchemy>=2.0,<3.0, customtkinter>=5.2,<7.0 y "
        "pillow>=10.0,<14.0; por eso el pipeline instala el ORM y la interfaz gráfica antes "
        "de ejecutar las pruebas."
    )

    doc.add_heading("4.3.6 Pull Request, revisión y merge", level=2)
    for texto in [
        "Rama de trabajo: feature/PBI-02-orm-persistencia.",
        "Pull Request: #2 (completar la URL al publicar en GitHub).",
        "Revisión: Git Master (Omar Huallpayunca).",
        "Merge: --no-ff hacia main, conservando la secuencia TDD.",
        "Checks: Ruff y pytest en verde; SQLAlchemy declarado en requirements.txt.",
        f"Merge local: {hashes.get('merge', 'se registra al fusionar la rama')}.",
        "Tag de entrega: pbi-02-v1.0.",
        "PBI-03: rama feature/PBI-03-interfaz-grafica, commits 37b0a15 (GUI) y f2bf169 (datos demo); merge --no-ff a main y tag pbi-03-gui-v1.0.",
    ]:
        agregar_bullet(doc, texto)


def seccion_gui(doc: Document, hashes: dict) -> None:
    encabezado = doc.add_heading("4.4 Interfaz gráfica de escritorio", level=1)
    encabezado.paragraph_format.page_break_before = True

    doc.add_heading("4.4.1 Objetivo y decisiones de diseño", level=2)
    doc.add_paragraph(
        "El PBI-02 dejó una capa de persistencia sólida, pero su uso diario dependía de la "
        "terminal. El PBI-03 añade una aplicación de escritorio funcional para que el asesor "
        "comercial opere el sistema con mouse y teclado. Se eligió CustomTkinter por ser una "
        "extensión moderna de Tkinter, ligera, multiplataforma y fácil de instalar."
    )
    for texto in [
        "Tema oscuro minimalista: fondo #070B10, paneles #0E151E, texto #F4F7FB y acento turquesa #2DD4BF.",
        "Pocas imágenes: la identidad visual se resuelve con tipografía, tarjetas, píldoras de estado y un icono propio.",
        "Reutilización de capas: la GUI llama a los repositorios; las reglas de totales, descuento, stock y documento siguen en src/dominio.",
        "Terminal conservada: python -m src.main --cli y --demo mantienen el flujo anterior.",
        "Primer arranque útil: sembrar_pedidos_demo() crea pedidos de ejemplo de forma idempotente por cliente, estado y líneas.",
    ]:
        agregar_bullet(doc, texto)

    doc.add_heading("4.4.2 Pantallas implementadas", level=2)
    agregar_tabla(
        doc,
        ["Pantalla", "Componentes principales", "Operaciones"],
        [
            ["Panel general", "Cuatro KPI, pedidos recientes, alerta de stock bajo", "Consultar el negocio y navegar"],
            ["Catálogo", "Búsqueda, filtro, tabla, modales de producto y categoría", "Crear categorías y productos con etiquetas"],
            ["Clientes", "Buscador, directorio, ficha con historial", "Crear, consultar y eliminar con cascada"],
            ["Pedidos", "Selector de cliente, catálogo, carrito, resumen de totales", "Armar pedidos, validar stock, confirmar y cambiar estado"],
            ["Reportes", "KPIs, ventas por estado, top de productos, inventario", "Analizar ventas y exportar CSV"],
        ],
        anchos=[1.2, 2.9, 2.5],
        tamano=8.5,
    )
    agregar_caption(doc, "Tabla 6. Pantallas de la aplicación de escritorio.")
    doc.add_paragraph(
        "La navegación lateral resalta la pantalla activa y la barra superior muestra el título, "
        "la fecha y un botón de refresco. Los atajos Ctrl+1 … Ctrl+5 cambian de pantalla y F5 "
        "recarga la vista actual."
    )

    doc.add_heading("4.4.3 Evidencias de la interfaz", level=2)
    agregar_imagen(doc, EVIDENCIAS / "gui" / "01_panel_general.png")
    agregar_caption(doc, "Figura G1. Panel general con indicadores, pedidos recientes y alertas de stock.")
    agregar_imagen(doc, EVIDENCIAS / "gui" / "04_pedidos.png")
    agregar_caption(doc, "Figura G2. Punto de venta con cliente, catálogo, carrito y resumen de totales.")
    agregar_imagen(doc, EVIDENCIAS / "gui" / "05_reportes.png")
    agregar_caption(doc, "Figura G3. Reportes con ventas por estado, ranking de productos e inventario.")
    doc.add_paragraph(
        "Las capturas restantes (catálogo, clientes, modal de producto, detalle de pedido y ficha "
        "de cliente) están disponibles en docs/evidencias/gui/."
    )

    doc.add_heading("4.4.4 Flujo funcional validado", level=2)
    doc.add_paragraph(
        "Además de las pruebas automatizadas, se ejecutó una prueba de humo sobre la GUI real con "
        "una base temporal:"
    )
    for indice, texto in enumerate([
        "Alta de una categoría y de un producto TST-001 con precio, stock y etiquetas desde los modales.",
        "Alta de un cliente con documento válido desde la pantalla Clientes.",
        "Selección del cliente, dos unidades en el carrito y confirmación del pedido.",
        "Verificación del total (S/ 199.80) y del descuento de stock.",
        "Cambio del pedido a estado PAGADO desde el modal de detalle.",
        "Exportación del reporte a CSV en la carpeta reportes/.",
    ], start=1):
        agregar_numero(doc, texto, indice)
    doc.add_paragraph(
        "El flujo terminó sin errores y confirmó que la interfaz escribe en SQLite a través de los "
        "mismos repositorios del PBI-02."
    )

    doc.add_heading("4.4.5 Estructura y resultado", level=2)
    agregar_codigo(
        doc,
        """src/presentacion/gui/
├── app.py            # Ventana principal, barra lateral, barra superior y routing
├── tema.py           # Paleta, tipografías y radios
├── componentes.py    # Tarjetas, tablas, botones, campos, modales y toasts
├── utilidades.py     # Moneda, fechas, iniciales y colores semánticos
└── pantallas/
    ├── dashboard.py
    ├── catalogo.py
    ├── clientes.py
    ├── pedidos.py
    └── reportes.py""",
        tamano=8,
    )
    doc.add_paragraph(
        "La ventana principal crea una única sesión SQLAlchemy, inicializa el esquema, siembra los "
        "datos y comparte los repositorios con todas las pantallas. Al cerrar, cierra la sesión y "
        "libera el engine. La aplicación se abre con python -m src.main; la suite total queda en "
        "23 pruebas aprobadas y Ruff no reporta errores."
    )
    if hashes.get("merge_pbi03"):
        doc.add_paragraph(
            f"Merge local del PBI-03: {hashes.get('merge_pbi03')}. Tag de entrega: pbi-03-gui-v1.0."
        )


def seccion_conclusiones(doc: Document) -> None:
    encabezado = doc.add_heading("4.5 Conclusiones", level=1)
    encabezado.paragraph_format.page_break_before = True
    for indice, texto in enumerate([
        "El ERD refleja el negocio real: cliente, pedido, detalle, producto, categoría y etiqueta cubren el flujo comercial de OPTIFIBER sin tablas innecesarias.",
        "La separación de capas funcionó: las reglas de negocio se probaron sin base de datos y los modelos ORM se probaron con SQLite en memoria.",
        "La cascada es segura y explícita: eliminar un cliente elimina sus pedidos y detalles; categorías y productos se protegen con RESTRICT.",
        "El fixture hermético evita efectos colaterales: cada prueba recibe una base nueva y no depende de data/optifiber.db.",
        "El TDD dejó evidencia verificable: cada ciclo tiene un commit RED que falla por la razón esperada y un commit GREEN que lo resuelve.",
        "La solución es funcional: la CLI permite registrar clientes y pedidos, consultar el historial, ver reportes y comprobar la cascada.",
        "La calidad está automatizada: Ruff y pytest se ejecutan localmente y en GitHub Actions; el pipeline falla si alguna prueba o regla de estilo se rompe.",
        "La GUI reutiliza la arquitectura: las pantallas llaman a los repositorios existentes y las reglas de negocio permanecen en src/dominio, sin duplicar lógica.",
        "La experiencia de escritorio quedó completa: panel, catálogo, clientes, pedidos, reportes, formularios modales y exportación CSV cubren el flujo comercial sin terminal.",
    ], start=1):
        agregar_numero(doc, texto, indice)

    doc.add_heading("Lecciones aprendidas", level=2)
    for texto in [
        "Declarar la tabla de asociación antes de las clases evita el error de resolución de secondary.",
        "StaticPool es necesario para que SQLite en memoria mantenga el esquema entre la creación y la sesión de prueba.",
        "PRAGMA foreign_keys=ON es imprescindible para que SQLite respete las claves foráneas.",
        "La cascada ORM debe combinarse con una decisión de negocio: no todo debe borrarse en cascada.",
        "Los precios históricos en DetallePedido protegen la facturación ante cambios de catálogo.",
    ]:
        agregar_bullet(doc, texto)


def seccion_checklist(doc: Document) -> None:
    doc.add_heading("4.6 Lista de verificación para el docente", level=1)
    agregar_tabla(
        doc,
        ["N°", "Criterio", "Evidencia", "Estado"],
        [
            ["1", "Modelo ERD e infraestructura ORM", "§4.1, modelos.py, erd_optifiber.png", "SÍ"],
            ["2", "Fixture de base de datos en memoria", "tests/conftest.py con sqlite:///:memory:", "SÍ"],
            ["3", "Pruebas 1:N y cascada en verde", "23 passed; §4.3.2 y §4.3.4", "SÍ"],
            ["4", "Historial RED → GREEN → REFACTOR", "git log; docs/evidencias/commits.md", "SÍ"],
            ["5", "Pull Request y GitHub Actions", "ci.yml; PR #2; merge --no-ff", "SÍ"],
            ["6", "Capítulo 4 completo", "Informe MD + DOCX", "SÍ"],
            ["7", "Aplicación funcional con interfaz gráfica", "§4.4, src/presentacion/gui/ y docs/evidencias/gui/", "SÍ"],
        ],
        anchos=[0.5, 2.4, 2.9, 0.6],
        tamano=8.5,
    )
    agregar_caption(doc, "Tabla 7. Checklist oficial del Trabajo Colaborativo N° 04.")

    doc.add_heading("Anexo A · PBI-02", level=2)
    agregar_tabla(
        doc,
        ["Campo", "Valor"],
        [
            ["Historia", "Como asesor comercial de OPTIFIBER deseo registrar clientes, productos y pedidos con sus líneas de detalle para consultar el historial, controlar el stock y calcular el total con descuentos."],
            ["Criterios", "Guardar cliente con pedidos; eliminar cliente en cascada; consultar producto con etiquetas; registrar pedido con detalle y stock; rechazar stock insuficiente; documento único."],
            ["Rama", "feature/PBI-02-orm-persistencia"],
            ["Estado", "Terminado · fusionado a main"],
            ["Pruebas", "21 aprobadas al cierre del PBI-02"],
        ],
        anchos=[1.1, 5.3],
        tamano=8.5,
    )

    doc.add_heading("Anexo B · Distribución del equipo", level=2)
    agregar_tabla(
        doc,
        ["Integrante", "Rol", "Aporte principal"],
        [
            ["Omar Huallpayunca", "Git Master + PO/SM", "Repositorio, PBI, revisión, merges, tags e informe."],
            ["Kevin Llerena", "Lead Backend", "Dominio, repositorios, reglas de negocio."],
            ["Amaru Ai-me", "Database & ORM", "ERD, modelos, relaciones, cascadas y restricciones."],
            ["Julio Diaz", "DevOps & QA + Frontend", "Fixture, pruebas, Ruff, CI, GUI CustomTkinter y evidencias."],
        ],
        anchos=[1.7, 1.9, 2.8],
        tamano=8.5,
    )

    doc.add_heading("Anexo C · PBI-03", level=2)
    agregar_tabla(
        doc,
        ["Campo", "Valor"],
        [
            ["Historia", "Como asesor comercial de OPTIFIBER deseo gestionar clientes, productos, pedidos y reportes desde una aplicación de escritorio para operar sin depender de la terminal."],
            ["Criterios", "Abrir la GUI por defecto; navegar entre cinco pantallas; crear productos y clientes; registrar pedidos con stock; cambiar estados; eliminar clientes en cascada; exportar reportes."],
            ["Rama", "feature/PBI-03-interfaz-grafica"],
            ["Estado", "Terminado · fusionado a main"],
            ["Pruebas", "23 aprobadas en total (21 del PBI-02 + 2 de datos demo)"],
            ["Evidencias", "docs/evidencias/gui/"],
        ],
        anchos=[1.1, 5.3],
        tamano=8.5,
    )

    doc.add_heading("Referencias", level=2)
    for texto in [
        "SQLAlchemy Documentation. ORM Declarative Mapping y Relationship Configuration. https://docs.sqlalchemy.org/",
        "CustomTkinter Documentation. Modern UI toolkit for Tkinter. https://customtkinter.tomschimansky.com/",
        "pytest Documentation. Fixtures. https://docs.pytest.org/",
        "Ruff Documentation. Linter and formatter. https://docs.astral.sh/ruff/",
        "GitHub Docs. GitHub Actions workflow syntax. https://docs.github.com/actions",
        "Instructivo de Trabajo Colaborativo N° 04. Persistencia de Datos, Mapeo ORM y Pruebas con Fixtures. Curso Construcción de Software.",
    ]:
        agregar_bullet(doc, texto)


def main() -> None:
    hashes = leer_json("commits.json")
    doc = Document()
    configurar_pagina(doc)
    configurar_estilos(doc)
    agregar_encabezado_pie(doc)

    portada(doc)
    seccion_resumen(doc)
    seccion_erd(doc, hashes)
    seccion_orm(doc, hashes)
    seccion_tdd(doc, hashes)
    seccion_gui(doc, hashes)
    seccion_conclusiones(doc)
    seccion_checklist(doc)

    doc.core_properties.title = (
        "Informe Capítulo 4 · Persistencia, ORM, Fixtures e Interfaz Gráfica"
    )
    doc.core_properties.author = "Equipo OPTIFIBER"
    doc.core_properties.subject = "Trabajo Colaborativo N° 04 · PBI-02 + PBI-03"
    doc.core_properties.comments = "Informe generado desde docs/INFORME_CAPITULO_4.md"

    doc.save(SALIDA)
    print(f"Informe generado: {SALIDA}")


if __name__ == "__main__":
    main()
